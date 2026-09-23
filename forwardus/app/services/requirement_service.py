"""수출요건 확인과 증빙 서류 업로드·분석."""

from __future__ import annotations

import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from app.collectors import ai_client
from app.models.requirement_document import RequirementDocument
from app.processors import export_requirements
from app.repositories import shipment_repository
from app.services import ServiceError
from app.validators import ValidationError
from app.validators.cargo_validator import PACKAGE_UNITS

# 올릴 수 있는 파일. 증명서는 대개 PDF나 스캔 이미지로 옵니다.
ALLOWED_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".txt", ".docx"}
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
UPLOAD_DIR = Path("instance") / "requirement_uploads"


def official_requirements(hs_code: str) -> dict:
    """관세청에 이 품목의 수출 요건을 직접 물어봅니다.

    우리 규칙표는 HS 류(類)로 짐작한 안내이고, 이쪽은 세번 10자리 기준의
    관세청 자료입니다. 답이 오면 그것이 기준이고, 우리 규칙은 보조가 됩니다.
    """

    from app.collectors import customs_extra_client

    result = customs_extra_client.export_requirement_laws(hs_code)
    if result["success"]:
        return {"available": True, "laws": result["data"], "source": "api",
                "note": ("관세법 제226조 세관장확인대상입니다. 아래 법령의 요건승인을 "
                         "받아야 수출신고가 수리됩니다." if result["data"] else
                         "이 세번은 세관장확인대상이 아닙니다. 다만 개별법 요건은 "
                         "따로 확인해야 합니다.")}
    # 키가 없거나 관세청이 멈춘 경우. 우리 규칙표로만 안내하고 그렇다고 밝힙니다.
    return {"available": False, "laws": [], "source": result["source"],
            "message": result["message"]}


def requirements_for(shipment) -> dict:
    """이 건에서 확인해야 할 수출 요건."""

    cargos = list(shipment.cargos)
    dangerous = any(cargo.is_dangerous for cargo in cargos)

    by_item = []
    seen = {}
    for cargo in cargos:
        items = export_requirements.check(cargo.hs_code, is_dangerous=cargo.is_dangerous)
        official = official_requirements(cargo.hs_code) if cargo.hs_code else {
            "available": False, "laws": [], "source": "", "message": "HS부호가 없습니다."}
        by_item.append({
            "line_no": cargo.line_no,
            "product_description": cargo.product_description,
            "hs_code": cargo.hs_code,
            "items": items,
            "official": official,
            "summary": export_requirements.summary(cargo.hs_code, items),
        })
        for item in items:
            seen.setdefault(item["key"], item)

    # 원산지증명서는 품목이 아니라 바이어와 협정이 정합니다. 따로 붙입니다.
    origin = export_requirements._as_item(export_requirements.ORIGIN_RULE)
    seen.setdefault(origin["key"], origin)

    combined = list(seen.values())
    uploads = list(shipment.requirement_documents)
    uploaded_keys = {doc.requirement_key for doc in uploads}
    for item in combined:
        item["uploaded"] = item["key"] in uploaded_keys

    return {
        "by_item": by_item,
        "requirements": combined,
        "dangerous": dangerous,
        "uploads": uploads,
        "ai_available": ai_client.available(),
        "lookup_links": export_requirements.LOOKUP_LINKS,
        "note": "여기 나오는 것은 '확인해야 할 것'입니다. 최종 판단은 세관과 수입국이 합니다.",
    }


def origin_certificates(shipment) -> list[RequirementDocument]:
    """이 건에 등록한 원산지증명서."""

    return [doc for doc in shipment.requirement_documents if doc.requirement_key == "origin"]


def _upload_root() -> Path:
    root = Path.cwd() / UPLOAD_DIR
    root.mkdir(parents=True, exist_ok=True)
    return root


def upload(shipment, file_storage, requirement_key: str,
           agreement: str = "") -> RequirementDocument:
    """증빙 서류를 저장합니다. 분석은 저장한 뒤 따로 부릅니다."""

    if file_storage is None or not (file_storage.filename or "").strip():
        raise ValidationError("올릴 파일을 선택해주세요.", "file")

    name = Path(file_storage.filename).name
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        allowed = ", ".join(sorted(ALLOWED_SUFFIXES))
        raise ValidationError(f"올릴 수 있는 형식은 {allowed} 입니다.", "file")

    data = file_storage.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValidationError(
            f"파일은 {MAX_UPLOAD_BYTES // (1024 * 1024)}MB까지 올릴 수 있습니다.", "file")
    if not data:
        raise ValidationError("빈 파일입니다.", "file")

    stored_name = f"{shipment.shipment_id}_{uuid.uuid4().hex}{suffix}"
    (_upload_root() / stored_name).write_bytes(data)

    titles = {item["key"]: item["title"] for item in requirements_for(shipment)["requirements"]}
    document = RequirementDocument(
        shipment=shipment,
        requirement_key=requirement_key or "",
        requirement_title=titles.get(requirement_key, "기타 증빙"),
        filename=name[:300],
        stored_name=stored_name,
        content_type=(file_storage.mimetype or "")[:120],
        size_bytes=len(data),
        agreement=str(agreement or "").strip()[:120],
    )
    shipment_repository.add(document)
    shipment_repository.commit()
    return document


def delete_upload(shipment, document_id: int) -> None:
    document = _get_upload(shipment, document_id)
    path = _upload_root() / document.stored_name
    path.unlink(missing_ok=True)
    shipment_repository.delete(document)
    shipment_repository.commit()


def _get_upload(shipment, document_id: int) -> RequirementDocument:
    document = next((doc for doc in shipment.requirement_documents if doc.id == document_id), None)
    if document is None:
        raise ServiceError("올린 서류를 찾을 수 없습니다.", "NOT_FOUND")
    return document


def requirement_key_of(shipment, document_id: int) -> str:
    return _get_upload(shipment, document_id).requirement_key


# 스캔 PDF를 OCR로 읽을 때 앞 몇 장만 봅니다. 증명서는 보통 한두 장이고, 장마다 몇 초 걸립니다.
OCR_PDF_PAGES = 3
# 글자가 이보다 적은 PDF는 스캔본으로 보고 OCR로 읽습니다.
SCANNED_TEXT_CHARS = 40


def extract_text(path: Path, suffix: str) -> str:
    """서류에서 글자를 뽑습니다. 못 뽑으면 빈 글자를 돌려줍니다.

    사진(PNG·JPG)과 글자가 없는 스캔 PDF는 Tesseract OCR로 읽습니다. (app/processors/ocr.py)
    """

    from app.processors import ocr

    try:
        if suffix == ".txt":
            return path.read_text(encoding="utf-8", errors="replace")
        if suffix in (".png", ".jpg", ".jpeg"):
            from PIL import Image

            with Image.open(path) as image:
                image.load()
                return ocr.read_text(image)
        if suffix == ".pdf":
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                # 증명서는 보통 한두 장입니다. 앞 5장이면 충분합니다.
                text = "\n".join((page.extract_text() or "") for page in pdf.pages[:5])
                if len(text.strip()) < SCANNED_TEXT_CHARS and ocr.available():
                    scanned = [ocr.read_text(page.to_image(resolution=200).original)
                               for page in pdf.pages[:OCR_PDF_PAGES]]
                    text = "\n\n".join(page for page in scanned if page.strip()) or text
                return text
        if suffix == ".docx":
            import docx

            return "\n".join(paragraph.text for paragraph in docx.Document(path).paragraphs)
    except Exception:
        # 어떤 형식이든 읽기에 실패하면 "못 읽었다"로 넘깁니다. 화면에서 이유를 알려 줍니다.
        return ""
    return ""


def _unreadable_message(suffix: str) -> str:
    """글자를 못 읽었을 때의 안내. OCR이 없어서인지, 있는데도 못 읽었는지 나눠 말합니다."""

    from app.processors import ocr

    scanned = suffix in (".png", ".jpg", ".jpeg", ".pdf")
    if scanned and not ocr.available():
        return ("서류에서 글자를 읽지 못했습니다. 사진·스캔 서류를 읽는 OCR(Tesseract, 한글·영문)이 "
                "이 서버에 설치되어 있지 않습니다. 글자가 있는 PDF로 올리거나 내용을 직접 확인해 주세요.")
    if scanned:
        return ("서류에서 글자를 읽지 못했습니다. 사진이 흐리거나 기울어져 있을 수 있습니다. "
                "밝은 곳에서 정면으로 다시 찍거나, 글자가 있는 PDF로 올려 주세요.")
    return "서류에서 글자를 읽지 못했습니다. 파일이 비어 있거나 손상되었을 수 있습니다."


def shipment_context(shipment) -> dict:
    """AI에게 넘길 이 건의 요약. 대조 기준이 되는 값만 넣습니다."""

    buyer = shipment.buyer
    return {
        "수출자": shipment.exporter_name,
        "수출자_사업자등록번호": shipment.exporter_business_no or "(미입력)",
        "구매자": buyer.name if buyer else "",
        "목적국": f"{shipment.destination_name} ({shipment.destination_country})",
        "출항예정일": shipment.etd.isoformat() if shipment.etd else "",
        "원산지": shipment.country_of_origin,
        "품목": [
            {
                "품명": cargo.product_description,
                "HS부호": cargo.hs_code,
                "수량": f"{cargo.quantity} {PACKAGE_UNITS.get(cargo.package_type, cargo.package_type)}",
                "총중량_kg": cargo.total_weight_kg,
                "위험물": (f"{cargo.un_number} CLASS {cargo.dg_class}"
                           if cargo.is_dangerous and cargo.un_number else ""),
            }
            for cargo in shipment.cargos
        ],
    }


def analyze(shipment, document_id: int) -> RequirementDocument:
    """올린 서류를 읽어 이 건과 대조합니다."""

    document = _get_upload(shipment, document_id)
    path = _upload_root() / document.stored_name
    if not path.exists():
        document.review_status = "failed"
        document.review_summary = "저장된 파일을 찾지 못했습니다. 다시 올려 주세요."
        shipment_repository.commit()
        return document

    suffix = Path(document.filename).suffix.lower()
    text = extract_text(path, suffix)
    if not text.strip():
        document.review_status = "failed"
        document.review_summary = _unreadable_message(suffix)
        document.review_findings = []
        document.reviewed_at = datetime.now(timezone.utc)
        shipment_repository.commit()
        return document

    context = shipment_context(shipment)
    context["확인하려는_요건"] = document.requirement_title
    if document.agreement:
        context["적용하려는_협정"] = document.agreement
    result = ai_client.review_document(_clean(text), context)

    if not result["success"]:
        document.review_status = "failed"
        document.review_summary = result["message"]
        document.review_findings = []
    else:
        data = result["data"]
        document.review_status = data["status"]
        document.review_summary = (f"{data['document_type']} · {data['summary']}"
                                   if data["document_type"] else data["summary"])
        document.review_findings = data["findings"]
    document.reviewed_at = datetime.now(timezone.utc)
    shipment_repository.commit()
    return document


def _clean(text: str) -> str:
    """빈 줄과 겹친 공백을 줄여 보냅니다."""

    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()
