"""사용성 점검 2회차(2026-10-04) — 서버 쪽.

  어디까지 읽었는지 말없이 버렸습니다(PDF 60쪽·스캔 10쪽), 메모장 'Unicode' 저장을 거절했고,
  국문 OCR 은 글자마다 띄어 읽어 '독소 0개'였습니다. 갑/을 표 머리·OCR 정의문을 못 읽었고,
  상담 창은 2,000자에서 잘린 계약서를 그대로 판정했습니다. 근거는 "U.S.A." 에서 잘렸습니다.
"""

from __future__ import annotations

import io

import pytest

from app import create_app, db
from app.processors import contract_clauses
from app.services import contract_clause_service as service
from config import TestConfig


def _pdf(pages: list[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for body in pages:
        doc.new_page().insert_text((50, 72), body, fontsize=10)
    return doc.tobytes()


def test_60쪽이_넘는_PDF_는_앞만_봤다고_알린다():
    pages = [f"SALES CONTRACT page {n}. General terms." for n in range(1, 63)]
    text, notes = service.read_contract("long.pdf", _pdf(pages))
    assert "page 60" in text and "page 61" not in text
    assert any("62쪽" in note and "60쪽" in note for note in notes)


def test_메모장_Unicode_저장도_읽는다():
    text = "수출계약서\n매수인은 재판매 대금을 회수한 후 지급한다."
    assert service.read_file("a.txt", text.encode("utf-16")) == text


def test_공백뿐인_파일은_빈_파일이다():
    with pytest.raises(Exception) as caught:
        service.read_contract("a.txt", b"   \r\n  ")
    assert "빈 파일" in str(caught.value)


def test_국문_OCR_의_글자마다_띄어쓰기를_붙인다():
    spaced = "갑 은 언 제 든 지 서 면 통 지 로 본 계약 을 해 지 할 수 있다"
    joined = service._join_hangul(spaced)
    assert "갑은" in joined and "해지할" in joined


def test_띄어쓰기가_멀쩡한_국문은_그대로다():
    line = "매수인은 언제든지 서면 통지로 본 계약을 해지할 수 있다"
    assert service._join_hangul(line) == line


@pytest.mark.parametrize("doc,label", [
    ("수출계약서\n구분 갑 (매도인) 을 (매수인) 상호 주식회사 한빛상사 ABC Trading\n"
     "제5조 을은 언제든지 서면 통지로 본 계약을 해지할 수 있다.", "갑"),
    ('물품공급 계약서 매수인 베트남 하노이 소재 VINA TRADING )5(( 이 하 " 갑 " 이 라 한 다 ) 와 매도인 '
     '경기도 안산시 소재 주식회사대한정밀 (이하 " 을 " 이라한다 ) 은 다음과 같이 계약한다.', "을"),
    ('수출계약서\n주식회사 한빛상사(이하 「갑」이라 한다)와 ABC Trading Co., Ltd.(이하 「을」이라 한다)는 '
     "다음과 같이 계약한다. 갑은 매도인으로서 물품을 공급한다.", "갑"),
])
def test_갑을_여러_서식에서_우리_쪽을_읽는다(doc, label):
    side = contract_clauses.our_side(doc)
    assert side and side["label"] == label


def _chat(question):
    app = create_app(TestConfig)
    with app.app_context():
        db.create_all()
        try:
            return app.test_client().post("/api/support-chat", json={"question": question}).get_json()
        finally:
            db.session.remove(); db.drop_all()


def test_상담_창은_잘린_계약서를_판정하지_않고_계약서_창으로_보낸다():
    long = "SALES CONTRACT\n" + ("1. The Seller shall deliver the goods as agreed in Annex 1. " * 40)
    for text in (long[:2000], long):
        answer = _chat(text)
        assert answer["source"] == "contract"
        assert "판정하지 않았습니다" in answer["data"]["answer"]
        assert answer["data"]["links"][0]["url"] == "/#contract-check"


def test_근거는_USA_에서_끊기지_않고_자르면_말줄임표를_단다():
    doc = ("SALES CONTRACT\nThis Contract shall be governed by the laws of the State of California, "
           "U.S.A. without regard to its conflict of laws rules.")
    evidence = contract_clauses.analyze(doc)["clauses"]["governing_law"]["evidence"]
    assert "U.S.A. without regard" in evidence
    long = ("SALES CONTRACT\nNeither party shall be liable for any failure caused by fire, flood, "
            "storm, earthquake, war, riot, strike, lockout, epidemic or pandemic, governmental action, "
            "embargo, port congestion, shortage of vessels or containers, failure of suppliers, power "
            "failure, cyber attack or any other cause beyond its reasonable control which is force majeure.")
    evidence = contract_clauses.analyze(long)["clauses"]["force_majeure"]["evidence"]
    assert evidence.endswith("…") or evidence.startswith("…") or len(evidence) < 250


def test_내려받는_문안은_안내를_상자_밖에_두고_조_번호를_뗀다():
    import docx

    document = docx.Document(io.BytesIO(service.clause_docx(["arbitration"])))
    box = [p.text for p in document.tables[0].cell(0, 0).paragraphs]
    assert box[0] == "ARBITRATION"
    assert not any(line.startswith("(") for line in box)
    plain = service.clause_plain(["lc_conformity"])
    assert "[영문 문안] (넣을 문구의 예)" in plain


# ── 사용성 2회차에서 '독소 0개'로 안심시키던 문장 ───────────────────────────────
@pytest.mark.parametrize("key,line", [
    ("payment_on_resale", "The Buyer shall pay within 60 days after the end of the month in which it "
                          "resells the Goods."),
    ("payment_on_resale", "제3조 매수인은 물품이 판매된 달의 말일로부터 60일 이내에 대금을 지급한다."),
    ("full_return", "The Buyer may return any unsold goods at the Seller's expense."),
    ("open_warranty", "The Seller warrants the Goods against defects for an unlimited period."),
])
def test_흔한_독소_문장을_놓치지_않는다(key, line):
    assert key in contract_clauses.find_in("SALES CONTRACT\n" + line)


@pytest.mark.parametrize("key,line", [
    ("payment_on_resale", "The Buyer shall pay within 60 days after the end of the month in which the "
                          "Goods are shipped."),
    ("full_return", "Unsold goods remain the property of the Buyer and may not be returned."),
])
def test_비슷하지만_정상인_문장은_짚지_않는다(key, line):
    assert key not in contract_clauses.find_in("SALES CONTRACT\n" + line)


@pytest.mark.parametrize("head", ["PURCHASE ORDER TERMS AND CONDITIONS", "GENERAL TERMS AND CONDITIONS OF PURCHASE"])
def test_발주서_약관은_계약서로_본다(head):
    """실무에서는 바이어 발주서 약관이 곧 계약 조건입니다(사용성 2회차)."""

    doc = (f"{head}\n1. These terms govern all purchase orders. 2. The Buyer may terminate any order for "
           "convenience at any time without liability. 3. The Buyer may set off any amounts it claims "
           "against the price. 4. The courts of Texas shall have exclusive jurisdiction. "
           "5. Payment within 90 days after receipt of invoice. " * 3)
    assert service.looks_like_contract(doc)


def test_발주서_한_장은_여전히_계약서가_아니다():
    doc = ("PURCHASE ORDER\nPO No. 4500123. Item: brake pads, 20,000 sets, USD 3.20, CIF Haiphong. "
           "Payment by T/T. Shipment by 30 November 2026. " * 6)
    assert not service.looks_like_contract(doc)
