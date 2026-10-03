"""독소조항을 위험 종류 8묶음으로 (2026-10-03).

왜 이 파일이 있나
  독소조항이 55개가 되자 화면과 내려받은 문안에서 한 줄로 늘어서 읽히지
  않았습니다. 수출자가 먼저 묻는 것은 "돈을 못 받는가 · 얼마나 물리는가"라서,
  그 순서대로 묶어 보여 줍니다.

  새 독소조항을 넣고 묶음을 안 붙이면 화면에서 '기타'로 떨어집니다. 그래서
  빠짐없이 붙었는지 여기서 봅니다.
"""

from __future__ import annotations

import io

import docx
import pytest

from app.processors import contract_clauses
from app.services import contract_clause_service as service

TOXIC = [row["key"] for row in contract_clauses.CLAUSES if row["category"] == "toxic"]


def test_독소조항마다_묶음이_있다():
    missing = [key for key in TOXIC if key not in contract_clauses.TOXIC_GROUP]
    assert not missing, missing


def test_묶음표에는_독소조항만_있다():
    extra = set(contract_clauses.TOXIC_GROUP) - set(TOXIC)
    assert not extra, extra


def test_묶음은_여덟이고_모두_쓰인다():
    ids = [group_id for group_id, _ in contract_clauses.TOXIC_GROUPS]
    assert len(ids) == 8
    assert set(contract_clauses.TOXIC_GROUP.values()) == set(ids)


def test_대금을_못_받는_위험이_맨_앞이다():
    first = contract_clauses.TOXIC_GROUPS[0][1]
    assert "대금" in first
    assert contract_clauses.TOXIC_GROUP["lc_soft_clause"] == contract_clauses.TOXIC_GROUPS[0][0]


def test_점검표의_독소조항은_묶음_순서로_나온다():
    rows = service.checklist("CIF")["toxic"]
    order = [group_id for group_id, _ in contract_clauses.TOXIC_GROUPS]
    seen = [order.index(row["group"]) for row in rows]
    assert seen == sorted(seen)
    assert all(row["group_label"] for row in rows)


def test_도착국_조항은_묶음_안에서_앞에_온다():
    rows = service.checklist("CIF", country="CN")["toxic"]
    ip_rows = [row for row in rows if row["group"] == contract_clauses.TOXIC_GROUP["cn_tech_transfer"]]
    assert ip_rows[0]["for_country"]


# 고른 순서와 상관없이 묶음 순서로 냅니다 — 뒤 묶음(⑦)을 먼저 골라도.
PICKED = ["foreign_forum", "lc_soft_clause", "arbitration", "unlimited_damages"]


def test_Word_문안도_묶음_소제목으로_나뉜다():
    document = docx.Document(io.BytesIO(service.clause_docx(PICKED)))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    labels = dict(contract_clauses.TOXIC_GROUPS)
    pay = labels[contract_clauses.TOXIC_GROUP["lc_soft_clause"]]
    dispute = labels[contract_clauses.TOXIC_GROUP["foreign_forum"]]
    assert lines.index(pay) < lines.index(dispute)
    # 소제목 바로 뒤에 그 묶음의 조항이 옵니다.
    assert "소프트 조항" in lines[lines.index(pay) + 1]
    # 필수조항은 독소 묶음 뒤에, 소제목 없이.
    assert any("[필수]" in line for line in lines[lines.index(dispute):])


def test_텍스트_문안도_묶음_소제목으로_나뉜다():
    body = service.clause_plain(PICKED)
    labels = dict(contract_clauses.TOXIC_GROUPS)
    pay = labels[contract_clauses.TOXIC_GROUP["lc_soft_clause"]]
    damages = labels[contract_clauses.TOXIC_GROUP["unlimited_damages"]]
    dispute = labels[contract_clauses.TOXIC_GROUP["foreign_forum"]]
    assert body.index(pay) < body.index(damages) < body.index(dispute) < body.index("[필수]")


@pytest.mark.parametrize("build", [service.clause_docx, service.clause_plain, service.clause_text])
def test_묶음을_나눠도_고른_조항은_모두_나온다(build):
    out = build(PICKED)
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(out)).paragraphs) \
        if isinstance(out, bytes) else out
    for key in PICKED:
        assert contract_clauses.by_key(key)["title"] in text, key


def test_상담_답변도_묶음_소제목으로_나뉜다():
    doc = ("SALES CONTRACT\nThe Buyer may set off any amounts it claims against the price. "
           "The courts of Shanghai shall have exclusive jurisdiction.")
    result = service.review(doc, "FOB")
    result["summary"] = ""
    text = service.as_text(result)
    labels = dict(contract_clauses.TOXIC_GROUPS)
    pay = f"**{labels['payment']}**"
    dispute = f"**{labels['dispute']}**"
    assert text.index(pay) < text.index("상계") < text.index(dispute)
