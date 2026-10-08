"""사용성 점검 3회차(2026-10-04) — 실서버·Edge 로 써 보며 나온 것.

섞인 PDF(글자 6쪽 + 스캔 2쪽)의 스캔 쪽을 말없이 버린 것이 가장 컸습니다. 나머지는 같은
말을 두 번 하거나, 알릴 말을 화면 밖에 두거나, 키보드가 창 밖으로 새는 것입니다.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from app import create_app, db
from app.processors import contract_clauses
from app.services import contract_clause_service as service
from config import TestConfig

STATIC = Path(__file__).resolve().parents[1] / "app" / "static"


def _scan_page(doc, lines: str) -> None:
    from PIL import Image, ImageDraw, ImageFont

    try:
        font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        pytest.skip("시험용 글꼴이 없는 환경")
    image = Image.new("RGB", (1700, 600), "white")
    ImageDraw.Draw(image).text((60, 80), lines, fill="black", font=font, spacing=18)
    png = io.BytesIO()
    image.save(png, format="PNG")
    page = doc.new_page()
    page.insert_image(page.rect, stream=png.getvalue())


def test_글자_쪽과_스캔_쪽이_섞인_PDF_는_스캔_쪽도_읽는다():
    from app.processors import bank_redaction, ocr

    if not (ocr.available() and bank_redaction.ocr_available()):
        pytest.skip("OCR 이 설치되지 않은 환경")
    import pytest
    pymupdf = pytest.importorskip("pymupdf")

    doc = pymupdf.open()
    for n in range(1, 4):
        doc.new_page().insert_text((50, 72), f"SALES CONTRACT page {n}\nThe Seller shall deliver the Goods "
                                             f"as described in Annex {n} of this Contract.", fontsize=10)
    _scan_page(doc, "The courts of Shanghai shall have\nexclusive jurisdiction.")
    text, notes = service.read_contract("mixed.pdf", doc.tobytes())
    assert "page 3" in text and "exclusive" in text.lower()
    assert any("4쪽은 스캔" in note for note in notes)


def test_사진은_스캔본이라고_하지_않는다():
    source = (Path(service.__file__)).read_text(encoding="utf-8")
    assert '"스캔본" if name.endswith(".pdf") else "사진"' in source


def _chat(question):
    app = create_app(TestConfig)
    with app.app_context():
        db.create_all()
        try:
            return app.test_client().post("/api/support-chat", json={"question": question}).get_json()
        finally:
            db.session.remove(); db.drop_all()


def test_입력칸이_자르고_공백을_지운_계약서도_판정하지_않는다():
    long = "SALES CONTRACT\n" + ("1. The Seller shall deliver the goods as agreed in Annex 1. " * 40)
    answer = _chat(long[:2000].strip()[:1998])
    assert answer["source"] == "contract"


def test_상담_창의_우리_쪽_줄은_이름표_꼴이다():
    doc = ('본 계약은 주식회사 대한정밀(이하 "을", 매도인)과 VINA TRADING JSC(이하 "갑", 매수인) 사이에 '
           "체결한다. 갑은 언제든지 서면 통지로 이 계약을 해지할 수 있다.")
    text = service.as_text(service.review("수출 계약서\n" + doc, "FOB") | {"summary": ""})
    assert "우리 쪽(매도인·수출자): **을" in text
    assert "을 **을" not in text


def test_근거_끝에_다음_조_번호를_붙이지_않는다():
    doc = ("SALES CONTRACT\n12. GOVERNING LAW This Contract shall be governed by the laws of the State of "
           "California, U.S.A. 13. NOTICES All notices shall be in writing.")
    row = contract_clauses.analyze(doc)["clauses"]["governing_law"]
    assert not row["evidence"].rstrip("…").endswith("13")


def test_독소로_짚은_관할_문장을_직접_확인에_다시_내지_않는다():
    doc = ("SALES CONTRACT\n13. DISPUTE RESOLUTION All disputes arising out of or in connection with this "
           "Contract shall be finally settled by the state and federal courts located in San Francisco, "
           "California, and each party waives any objection to such venue. The parties agree to trial by jury, "
           "and punitive damages may be awarded.")
    result = service.review(doc, "FOB", "US")
    assert "us_jury_punitive" in {row["key"] for row in result["toxic"]}
    assert not [c for c in result["check"] if "San Francisco" in c["sentence"]]


def test_선적_근거는_결제_문장이_아니라_선적_문장이다():
    doc = ("SALES CONTRACT\nPayment: 30% advance, and the balance by T/T within 10 days after shipment. "
           "Shipment shall be made from Busan to Oakland not later than 30 Nov. 2026.")
    row = contract_clauses.analyze(doc)["clauses"]["shipment"]
    assert row["evidence"].startswith("Shipment shall be made")


def test_FOB_점검표에는_DDP_수입자_조항에_도착국_배지를_달지_않는다():
    rows = service.checklist("FOB", country="VN")["toxic"]
    assert not next(row for row in rows if row["key"] == "ddp_no_ior")["for_country"]
    rows = service.checklist("DDP", country="VN")["toxic"]
    assert next(row for row in rows if row["key"] == "ddp_no_ior")["for_country"]


# ── 화면 (JS·CSS) — 글자로 확인합니다 ────────────────────────────────────────
def test_접힌_묶음_안은_Tab_가두기에서_뺀다():
    js = (STATIC / "js" / "contract_modal.js").read_text(encoding="utf-8")
    assert "!folded(el)" in js and 'details:not([open])' in js


def test_Esc_한_번에_상담_창까지_닫지_않는다():
    modal = (STATIC / "js" / "contract_modal.js").read_text(encoding="utf-8")
    chat = (STATIC / "js" / "support_chat.js").read_text(encoding="utf-8")
    assert 'event.key === "Escape") { event.preventDefault()' in modal
    assert "event.defaultPrevented" in chat


def test_차분한_고치는_법과_갑을_경고는_덮이지_않는다():
    css = (STATIC / "css" / "contract.css").read_text(encoding="utf-8")
    assert ".cc_fix.cc_fix_calm" in css and ".cc_side.cc_side_warn" in css


def test_결과_위에_요약을_두고_뺀_조항을_모두_이름으로_알린다():
    js = (STATIC / "js" / "contract_clauses.js").read_text(encoding="utf-8")
    assert 'class="cc_summary"' in js
    assert "이 계약에 해당하지 않아 뺐습니다" in js
    assert "에는 해당하지 않아 고른" in js           # 인코텀즈를 바꿔 빠진 것
    assert "setTimeout" in js and "clearTimeout(slowTimer)" in js
    assert "data-cc-pickmsg" in js


# ── 3회차 마무리 (2026-10-04, '수정하고 다시 확인') ─────────────────────────────
def test_같은_근거_문장은_두_번째부터_앞_조항을_가리킨다():
    doc = ("SALES CONTRACT\nIn case of delay in shipment, the Seller shall pay liquidated damages of 2% of "
           "the contract price per day of delay, without any maximum limit.")
    result = service.review(doc, "FOB")
    same = [row for row in result["toxic"] if row["same_as"]]
    firsts = [row for row in result["toxic"] if not row["same_as"]]
    assert len(result["toxic"]) >= 2 and same and firsts
    assert "조항과 같은 문장" in service.as_text(result | {"summary": ""})


def test_다른_근거_문장은_가리키지_않는다():
    doc = ("SALES CONTRACT\nThe Buyer may set off any amount claimed against any sum payable to the Seller. "
           "This warranty shall survive indefinitely.")
    assert not [row for row in service.review(doc, "FOB")["toxic"] if row["same_as"]]


def test_거래_조건으로_짚은_독소는_까닭을_적는다():
    result = service.review("SALES CONTRACT\nDelivery DDP Sao Paulo.", "DDP", "BR")
    text = service.as_text(result | {"summary": ""})
    assert "까닭: 인코텀즈 DDP" in text


def test_전체_쪽보다_적게_올리면_나머지를_올리라고_한다():
    assert "2쪽이라고" in service._missing_pages("SALES CONTRACT ... Page 1 of 2", 1)


@pytest.mark.parametrize("text,pages", [("Page 1 of 1", 1), ("Page 2 of 2", 2), ("Annex 1. Price list", 1),
                                        ("USD 1/3 of the price", 1)])
def test_쪽을_다_올렸거나_쪽_표시가_없으면_알리지_않는다(text, pages):
    assert service._missing_pages(text, pages) == ""


def test_우리가_공제하는_권리는_바이어_상계가_아니다():
    doc = ("SALES CONTRACT\nThe Seller shall refund the Deposit to the Buyer (subject to the Seller's right to "
           "deduct from the Deposit any then-outstanding amounts owed by the Buyer).")
    assert "buyer_set_off" not in contract_clauses.find_in(doc)


def test_화면은_도착국_이름을_확인시키고_받기_단추를_겹쳐_보이지_않는다():
    js = (STATIC / "js" / "contract_clauses.js").read_text(encoding="utf-8")
    assert "Intl.DisplayNames" in js and "알 수 없는 나라 코드" in js
    assert "IntersectionObserver" in js and "actionsInView" in js
    assert 'lastSource === "text"' in js          # 4회차: 비우지 않고 마지막에 넣은 쪽으로 판정
    assert "조항과 같은 문장" in js


# ── 사용성 4회차 (2026-10-05) ──────────────────────────────────────────────────
FILLED = ("SALES CONTRACT\n1. GOODS Polypropylene 500 MT.\n2. PRICE FOB Busan, all prices in USD.\n"
          "3. PAYMENT T/T within 30 days after B/L date.\n4. SHIPMENT not later than 30 Nov. 2026.")


def test_빈칸이_그대로인_계약서는_다_갖췄다고_하지_않는다():
    blank = ("SALES CONTRACT\n1. PAYMENT <L/C at sight | T/T>.\n2. SHIPMENT Latest date of shipment: ________.\n"
             "3. PRICE FOB Busan, all prices are in [●].")
    result = service.review(blank, "FOB")
    assert result["blanks"] >= 3
    weak = {row["key"] for row in result["weak"]}
    assert {"shipment"} <= weak
    assert "빈칸" in service.as_text(result)


def test_채운_계약서와_수치_괄호_이메일은_빈칸이_아니다():
    doc = FILLED + "\nContact: <sales@hana.co.kr>. Late interest <3>% per month."
    assert service.review(doc, "FOB")["blanks"] == 0


def test_본문의_CIF_를_읽어_보험_누락을_짚는다():
    doc = ("SALES CONTRACT\n1. GOODS Steel coil 500 MT.\n2. PRICE CIF Rotterdam, USD 500 per MT.\n"
           "3. PAYMENT T/T within 30 days after B/L date.")
    result = service.review(doc, "")
    assert result["incoterms"] == "CIF" and result["incoterms_from_doc"] == "CIF"
    assert "insurance" in {row["key"] for row in result["missing"]}


def test_고른_인코텀즈와_본문이_다르면_알린다():
    doc = "SALES CONTRACT\n2. PRICE CIF Rotterdam, USD 500 per MT. PAYMENT T/T."
    result = service.review(doc, "DDP")
    assert result["incoterms_mismatch"] == "CIF"
    assert "CIF" in service.as_text(result) and "다시 판정" in service.as_text(result)


@pytest.mark.parametrize("text", ["hello", "안녕하세요 견적 부탁드립니다", "\x00" * 200])
def test_계약서가_아닌_글은_판정하지_않는다(text):
    with pytest.raises(service.ServiceError) as caught:
        service.review(text, "FOB")
    assert caught.value.error_code in {"NOT_CONTRACT", "VALIDATION_ERROR"}


def test_중국어_계약서는_독소_0개로_안심시키지_않는다():
    doc = "销售合同\n甲方可随时无理由解除本合同。乙方不得抵销任何款项。争议提交上海国际仲裁中心仲裁，适用中华人民共和国法律。" * 5
    with pytest.raises(service.ServiceError) as caught:
        service.review(doc, "FOB", "CN")
    assert caught.value.error_code == "UNSUPPORTED_LANG"


def test_도착국은_두_글자가_아니면_비운다():
    doc = FILLED
    assert service.review(doc, "FOB", "<B>X</B>" * 20)["country"] == ""
    assert service.review(doc, "FOB", "vn")["country"] == "VN"


def test_Word_변경_추적으로_넣은_문장도_읽는다():
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = docx.Document()
    document.add_paragraph("SALES CONTRACT")
    paragraph = document.add_paragraph("The Seller shall deliver the goods. ")
    ins = parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="Buyer" w:date="2026-10-01T00:00:00Z"><w:r>'
                    f'<w:t>The Buyer may terminate this Contract at any time for convenience.</w:t></w:r></w:ins>')
    paragraph._p.append(ins)
    deleted = parse_xml(f'<w:del {nsdecls("w")} w:id="2" w:author="Buyer" w:date="2026-10-01T00:00:00Z"><w:r>'
                        f'<w:delText>DELETED-SENTENCE</w:delText></w:r></w:del>')
    paragraph._p.append(deleted)
    section = document.sections[0]
    section.footer.paragraphs[0].text = "Footer: The Buyer may set off any amount claimed against the price."
    out = io.BytesIO()
    document.save(out)
    text, notes = service.read_contract("tracked.docx", out.getvalue())
    assert "terminate this Contract at any time" in text and "DELETED-SENTENCE" not in text
    assert "set off any amount" in text
    assert any("변경 추적" in note for note in notes)
    assert "termination_at_will" in contract_clauses.find_in(text)


def _client():
    app = create_app(TestConfig)
    ctx = app.app_context()
    ctx.push()
    db.create_all()
    return app.test_client(), ctx


def test_JSON_배열을_보내도_500이_아니다():
    client, ctx = _client()
    try:
        response = client.post("/contract/review", data="[1,2,3]", content_type="application/json")
        assert response.status_code == 400
        response = client.post("/contract/review", json={"text": ["x"]})
        assert response.status_code == 400
    finally:
        db.session.remove(); db.drop_all(); ctx.pop()


def test_받기는_겹친_조항을_빼고_상한을_둔다():
    client, ctx = _client()
    try:
        response = client.post("/contract/export", json={"keys": ["payment"] * 2000, "format": "txt"})
        body = response.get_data(as_text=True)
        assert response.status_code == 200 and body.count("[필수]") == 1
        assert response.headers["Content-Type"].count("charset") == 1
    finally:
        db.session.remove(); db.drop_all(); ctx.pop()


def test_받는_문안에_빈칸_설명과_설명_표시가_있다():
    text = service.clause_plain(["payment", "shipment"])
    assert "고쳐 넣을 칸" in text and "설명:" in text


def test_판정_결과를_파일로_받을_수_있다():
    result = service.review("SALES CONTRACT\nThe Buyer may set off any amount against the price.", "FOB")
    report = service.report_plain(result)
    assert report.startswith("﻿") and "\r\n" in report and "지우거나 고쳐야 할 조항" in report
    client, ctx = _client()
    try:
        data = client.post("/contract/review", json={"text": "SALES CONTRACT\nThe Buyer may set off any amount "
                                                              "against the price."}).get_json()["data"]
        assert "지우거나 고쳐야 할 조항" in data["report"]
    finally:
        db.session.remove(); db.drop_all(); ctx.pop()


def test_독소가_없어도_직접_확인_항목이_있으면_안심시키지_않는다():
    result = service.review(FILLED, "FOB", "SA")
    result["toxic"], result["missing"], result["weak"] = [], [], []
    result["watch_country"] = [{"key": "x", "title": "t", "risk": "r"}]
    assert "직접 확인하세요" in service.as_text(result)


def test_화면은_마지막에_넣은_쪽으로_판정하고_기준을_적는다():
    js = (STATIC / "js" / "contract_clauses.js").read_text(encoding="utf-8")
    assert "판정한 것:" in js and "판정 기준:" in js and "data-cc-report" in js
    assert 'class="cc_tag cc_watch">주의' in js and "ALIAS" in js and "ARE:" in js
    assert 'aria-label="계약서 본문 붙여 넣기"' in js and "getFullYear()" in js
