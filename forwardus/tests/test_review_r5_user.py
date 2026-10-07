"""전수 점검 5회차(사용자) — 4회차 수정이 다른 길에 남긴 구멍."""

import pytest

from app.collectors import location_client


@pytest.mark.parametrize("name, code", [("Vietnam", "VN"), ("USA", "US"), ("UK", "GB"), ("UAE", "AE"),
                                         ("Turkey", "TR"), ("영국", "GB"), ("Germany", "DE")])
def test_흔한_나라_표기를_코드로_찾는다(name, code):
    assert location_client.find_country_by_name(name) == code


def test_없는_나라는_찾지_않는다():
    assert location_client.find_country_by_name("Narnia") is None


@pytest.mark.parametrize("query", ["new york", "호치민", "도쿄", "NY"])
def test_항공_도착지를_도시_이름으로도_찾는다(query):
    assert location_client.search_locations(query, "airport").get("data")


def test_서류_작성_경로의_Buyer_국가도_잘라_쓰지_않는다(app):
    from app.services import document_pipeline_service as pipeline
    source = open(pipeline.__file__, encoding="utf-8").read()
    assert "upper()[:2]" not in source.split("buyer_country")[1][:400]


@pytest.mark.parametrize("name, code", [("Vietnam", "VN"), ("UK", "GB"), ("베트남", "VN")])
def test_서류_시작도_같은_표기를_받는다(name, code):
    from app.services.document_start_service import _country_code
    assert _country_code({"buyer_country": name}) == code


def test_중복_제출도_이동할_주소를_준다(app, client, shipment_payload):
    from app.services import planning_service

    payload = dict(shipment_payload)
    payload["schedule_id"] = planning_service.search_schedules(payload)["items"][0]["schedule_id"]
    first = client.post("/planning/api/shipments", json=payload).get_json()
    again = client.post("/planning/api/shipments", json=payload).get_json()
    assert first["success"] and again["success"]
    assert again["data"].get("reused") is True
    assert again["data"]["url"].startswith("/documents/")


def test_docx_표_안의_글자도_읽는다(tmp_path):
    import docx

    from app.services import requirement_service

    document = docx.Document()
    document.add_paragraph("CERTIFICATE OF ORIGIN")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Exporter"
    table.cell(0, 1).text = "FORWARDUS CO., LTD."
    table.cell(1, 0).text = "HS Code"
    table.cell(1, 1).text = "3306.10"
    path = tmp_path / "coo.docx"
    document.save(path)
    text = requirement_service.extract_text(path, ".docx")
    assert "FORWARDUS CO., LTD." in text and "3306.10" in text
