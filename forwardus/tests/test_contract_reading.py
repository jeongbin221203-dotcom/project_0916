"""계약서 파일 읽기와 당사자 읽기 — 사용성·전문가 점검에서 나온 버그 (2026-10-04).

  PDF 는 앞 5쪽, 스캔은 앞 3쪽만 읽었습니다(오퍼시트용 읽기 함수를 같이 써서).
  8쪽 계약서 끝의 중재 조항을 '빠짐'이라 하고, 뒤쪽 독소를 놓쳤습니다.
  메모장 기본 저장(CP949) .txt 는 글자가 깨진 채 '독소 0개'가 나왔습니다.
  쪽 번호만 글자로 박힌 스캔본은 OCR 을 하지 않고 거절했습니다.
  국문 갑/을 계약서는 방향을 못 봐서 우리 권리를 독소로 짚었습니다.
  Supplier/Distributor 로 부르는 계약서는 Seller/Buyer 규칙에 안 걸렸습니다.
  MAX_TEXT 를 넘는 글은 뒤를 안 보면서 '다 읽었다'고 했습니다.
"""

from __future__ import annotations

import io

import pytest

from app.processors import contract_clauses
from app.services import ServiceError
from app.services import contract_clause_service as service

KO = ("수출 매매계약서\n매도인 주식회사 한빛상사와 매수인 ABC Trading 은 다음과 같이 계약한다.\n"
      "제3조 매수인은 물품을 재판매하여 대금을 회수한 후 60일 이내에 대금을 지급한다.\n")


# ── 텍스트 파일 ─────────────────────────────────────────────────────────────
def test_메모장_기본_저장_CP949_를_읽는다():
    text = service.read_file("계약서.txt", KO.encode("cp949"))
    assert "재판매" in text
    assert "payment_on_resale" in contract_clauses.find_in(text)


def test_BOM_이_붙은_UTF8_도_읽는다():
    assert service.read_file("a.txt", KO.encode("utf-8-sig")).startswith("수출")


def test_글자가_아닌_파일은_깨진_채_판정하지_않는다():
    with pytest.raises(ServiceError) as caught:
        service.read_file("a.txt", bytes(range(256)) * 40)
    assert caught.value.error_code == "UNREADABLE"


def test_빈_파일은_빈_파일이라고_한다():
    with pytest.raises(ServiceError) as caught:
        service.read_file("a.txt", b"")
    assert "빈 파일" in str(caught.value)


# ── Word ────────────────────────────────────────────────────────────────────
def test_Word_계약서를_읽는다():
    import docx

    document = docx.Document()
    document.add_paragraph("SALES CONTRACT")
    document.add_paragraph("The courts of Shanghai shall have exclusive jurisdiction.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Payment"
    table.cell(0, 1).text = "The Buyer may set off any amounts it claims against the price."
    out = io.BytesIO()
    document.save(out)
    text = service.read_file("계약서.docx", out.getvalue())
    found = contract_clauses.find_in(text)
    assert {"foreign_forum", "buyer_set_off"} <= found


# ── PDF ─────────────────────────────────────────────────────────────────────
def _pdf(pages: list[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for body in pages:
        doc.new_page().insert_text((50, 72), body, fontsize=10)
    return doc.tobytes()


def test_PDF_는_뒤쪽까지_읽는다():
    pages = [f"SALES CONTRACT page {n}\nGeneral terms on page {n}." for n in range(1, 8)]
    pages.append("9. ARBITRATION All disputes shall be finally settled by arbitration in Seoul "
                 "under the rules of the Korean Commercial Arbitration Board.\n"
                 "The courts of Shanghai shall have exclusive jurisdiction.")
    text = service.read_file("long.pdf", _pdf(pages))
    assert "page 7" in text and "Arbitration Board" in text
    result = service.review(text, "FOB")
    assert "arbitration" not in {row["key"] for row in result["missing"]}
    assert "foreign_forum" in {row["key"] for row in result["toxic"]}


def test_쪽_번호만_글자인_스캔본은_OCR_로_읽는다():
    from app.processors import bank_redaction, ocr

    if not (ocr.available() and bank_redaction.ocr_available()):
        pytest.skip("OCR 이 설치되지 않은 환경")
    from PIL import Image, ImageDraw, ImageFont
    import pymupdf

    try:
        font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        pytest.skip("시험용 글꼴이 없는 환경")
    image = Image.new("RGB", (1700, 600), "white")
    ImageDraw.Draw(image).text(
        (60, 80), "SALES CONTRACT\nThe courts of Shanghai shall have\nexclusive jurisdiction.",
        fill="black", font=font, spacing=18)
    png = io.BytesIO()
    image.save(png, format="PNG")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=png.getvalue())
    page.insert_text((280, 820), "Page 1 of 1", fontsize=8)
    text = service.read_file("scan.pdf", doc.tobytes())
    assert "exclusive" in text.lower()


# ── 40만 자 ─────────────────────────────────────────────────────────────────
def test_한도를_넘는_글은_앞만_봤다고_밝힌다():
    # 같은 문장을 되풀이한 글은 반복 검사(전수 점검 4회차)가 거절하므로, 조항마다 다른 긴 글로 시험합니다.
    long = "SALES CONTRACT\n" + " ".join(
        f"Clause {n}: the parties agree on item {n * 7919 % 10007} at {n * 31 % 977} units per lot."
        for n in range(1, 7000))
    result = service.review(long, "FOB")
    assert result["truncated"] is True
    assert result["checked"] == service.MAX_TEXT
    assert "앞" in service.as_text(result | {"summary": ""})


def test_한도_안의_글은_잘렸다고_하지_않는다():
    assert service.review("SALES CONTRACT\nPayment by T/T.", "FOB")["truncated"] is False


# ── 갑/을 ───────────────────────────────────────────────────────────────────
GAB_SELLER = ('주식회사 대한산업(이하 "갑", 매도인)과 VINA Co., Ltd.(이하 "을", 매수인)는 '
              "다음과 같이 수출계약을 체결한다.\n")
GAB_BUYER = ('매수인 ABC Trading(이하 "갑")과 매도인 주식회사 한빛상사(이하 "을")는 '
             "다음과 같이 계약한다.\n")


def test_갑이_매도인이면_갑의_해지권은_독소가_아니다():
    doc = GAB_SELLER + "제5조 갑은 언제든지 서면 통지로 본 계약을 해지할 수 있다."
    assert "termination_at_will" not in contract_clauses.find_in(doc)


def test_을이_매수인이면_을의_해지권은_독소다():
    doc = GAB_SELLER + "제5조 을은 언제든지 서면 통지로 본 계약을 해지할 수 있다."
    assert "termination_at_will" in contract_clauses.find_in(doc)


def test_을이_매도인이면_갑의_무제한_배상은_우리에게_유리하다():
    doc = GAB_BUYER + "제6조 갑은 을의 모든 손해를 한도 없이 배상한다."
    assert "unlimited_damages" not in contract_clauses.find_in(doc)
    doc = GAB_BUYER + "제6조 을은 갑의 모든 손해를 한도 없이 배상한다."
    assert "unlimited_damages" in contract_clauses.find_in(doc)


def test_갑을_을_우리_쪽으로_읽었다고_밝힌다():
    side = contract_clauses.our_side(GAB_BUYER + "제1조 물품은 별지와 같다.")
    assert side and side["label"] == "을" and side["other_label"] == "갑"


def test_조사_을은_바꾸지_않는다():
    """'물품을' 의 을은 당사자가 아닙니다."""

    doc = GAB_SELLER + "제7조 을은 물품을 인수한 후 대금을 지급한다."
    body = contract_clauses.as_seller_with(contract_clauses.our_side(doc), doc)
    assert "물품을" in body and "매수인은" in body


def test_매수인이_배상하는_국문_문장은_무제한_배상이_아니다():
    doc = "수출계약서\n제6조 매수인은 매도인의 모든 손해를 한도 없이 배상한다."
    assert "unlimited_damages" not in contract_clauses.find_in(doc)


# ── Supplier / Distributor ──────────────────────────────────────────────────
DIST = ("DISTRIBUTION AGREEMENT between Hana Co., Ltd. (the \"Supplier\") and Euro GmbH "
        "(the \"Distributor\"). The Supplier shall supply the Products to the Distributor. "
        "The Distributor shall resell the Products in the Territory. ")


@pytest.mark.parametrize("key,line", [
    ("eu_gdpr_indemnity", "The Supplier shall indemnify and hold the Distributor harmless against "
                          "any fines or penalties imposed under the General Data Protection Regulation."),
    ("cert_test_cost", "All costs of certification, testing and factory audits required in the "
                       "Distributor's country shall be borne by the Supplier."),
    ("termination_at_will", "The Distributor may terminate this Agreement at any time for "
                            "convenience upon written notice."),
])
def test_Supplier_Distributor_계약서에서도_독소를_찾는다(key, line):
    assert key in contract_clauses.find_in(DIST + line)


def test_Distributor_가_무는_것은_독소가_아니다():
    line = ("The Distributor shall indemnify and hold the Supplier harmless against any fines "
            "imposed under the General Data Protection Regulation.")
    assert "eu_gdpr_indemnity" not in contract_clauses.find_in(DIST + line)


def test_Seller_Buyer_계약서의_supplier_는_바꾸지_않는다():
    """원자재 공급사(supplier)가 따로 나오는 매매계약 — 당사자가 아닙니다."""

    doc = ("SALES CONTRACT between the Seller and the Buyer. The Seller shall deliver the Goods. "
           "The Buyer shall pay the price. The Seller's raw material supplier shall be approved "
           "by the Buyer. The Buyer may inspect the Goods.")
    assert contract_clauses._as_parties(doc) == doc
