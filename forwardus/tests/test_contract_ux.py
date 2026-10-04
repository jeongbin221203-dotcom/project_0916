"""사용성 점검(2026-10-04)에서 나온 것 — 서버 쪽.

  상담 창: 기능을 꺼도 계약서 판정이 나왔고, 함께 붙는 링크(/documents)는 404 였습니다.
  Word 문안: <30>·<INCOTERMS> 같은 빈칸이 눈에 띄지 않았습니다.
  업로드: Word(.docx)를 받지 않았습니다.
"""

from __future__ import annotations

import io

from app import create_app, db
from app.services import contract_clause_service as service
from config import TestConfig

SOUND = ("SALES CONTRACT\n"
         "1. DESCRIPTION OF GOODS: The commodity, specification and quantity are set out in Annex 1.\n"
         "2. PRICE: FOB Busan, Incoterms 2020. The total contract value is USD 120,000.\n"
         "3. PAYMENT: by irrevocable L/C at sight to be opened at least 30 days before shipment.\n"
         "4. SHIPMENT: by 30 November 2026. Partial shipment and transhipment are allowed.\n"
         "5. INSPECTION: Inspection by SGS at the port of loading shall be final as to quality.\n"
         "6. FORCE MAJEURE: Neither party shall be liable for failure to perform due to force majeure.\n"
         "7. JURISDICTION: The courts of Shanghai shall have exclusive jurisdiction.\n"
         "8. SET-OFF: The Buyer may set off any amounts it claims against the price.\n"
         "9. GOVERNING LAW: This Contract shall be governed by the laws of the Republic of Korea.\n")


class OffConfig(TestConfig):
    CONTRACT_CLAUSES_ON = False


def _client(config):
    app = create_app(config)
    ctx = app.app_context()
    ctx.push()
    db.create_all()
    return app, ctx


def _support(client, text):
    return client.post("/api/support-chat", json={"question": text}).get_json()


def test_기능을_끄면_상담_창에서도_계약서로_판정하지_않는다():
    app, ctx = _client(OffConfig)
    try:
        answer = _support(app.test_client(), SOUND)
        assert (answer.get("source") or "") != "contract"
    finally:
        db.session.remove(); db.drop_all(); ctx.pop()


def test_상담_창_링크는_계약서_창을_연다():
    app, ctx = _client(TestConfig)
    try:
        answer = _support(app.test_client(), SOUND)
        assert answer["source"] == "contract"
        urls = [link["url"] for link in answer["data"]["links"]]
        assert urls == ["/#contract-check"]
        assert app.test_client().get("/").status_code in (200, 302)
    finally:
        db.session.remove(); db.drop_all(); ctx.pop()


def test_Word_문안의_빈칸은_노랗게_칠한다():
    import docx
    from docx.enum.text import WD_COLOR_INDEX

    document = docx.Document(io.BytesIO(service.clause_docx(["incoterms"])))
    box = document.tables[0].cell(0, 0)
    marked = [run.text for p in box.paragraphs for run in p.runs
              if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert "<INCOTERMS>" in marked


def test_창구가_Word_계약서를_받는다():
    import docx

    app, ctx = _client(TestConfig)
    try:
        document = docx.Document()
        document.add_paragraph(SOUND)
        out = io.BytesIO()
        document.save(out)
        out.seek(0)
        answer = app.test_client().post("/contract/review", data={
            "file": (out, "계약서.docx"), "incoterms": "FOB"}, content_type="multipart/form-data")
        data = answer.get_json()
        assert data["success"], data
        assert "buyer_set_off" in {row["key"] for row in data["data"]["toxic"]}
    finally:
        db.session.remove(); db.drop_all(); ctx.pop()
